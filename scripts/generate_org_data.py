#!/usr/bin/env python3
"""Generate org-chart.xlsx (source of truth) and org-chart.json (for the web app)."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
PHOTOS_DIR = ROOT / "assets" / "photos"

# Real ESS Claims Continuous Improvement roster (from the official chart).
# photo: filename in assets/photos/ OR a full URL. Leave blank for placeholder.
PEOPLE = [
    {
        "id": "1",
        "name": "Paulina Strozyk-Kozlowska",
        "role": "Head of ESS Claims",
        "manager_id": "",
        "team": "Leadership",
        "pillar": "leadership",
        "photo": "",
    },
    {
        "id": "2",
        "name": "Rebecca Newcastle",
        "role": "ESS Claims CI Lead",
        "manager_id": "1",
        "team": "Leadership",
        "pillar": "leadership",
        "photo": "",
    },
    {
        "id": "3",
        "name": "Deekshith Kiran",
        "role": "Pillar Co-Lead",
        "manager_id": "2",
        "team": "System Enhancement",
        "pillar": "system",
        "photo": "",
    },
    {
        "id": "4",
        "name": "Michele Panek",
        "role": "Pillar Lead",
        "manager_id": "2",
        "team": "System Enhancement",
        "pillar": "system",
        "photo": "",
    },
    {
        "id": "5",
        "name": "Vandana Taneja",
        "role": "Pillar Lead",
        "manager_id": "2",
        "team": "Process Improvement",
        "pillar": "process",
        "photo": "",
    },
    {
        "id": "6",
        "name": "Smriti Kandari",
        "role": "Pillar Lead",
        "manager_id": "2",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    # System Enhancement team
    {
        "id": "7",
        "name": "Anna Rzepka",
        "role": "System Enhancement",
        "manager_id": "4",
        "team": "System Enhancement",
        "pillar": "system",
        "photo": "",
    },
    {
        "id": "8",
        "name": "Sudha S.",
        "role": "System Enhancement",
        "manager_id": "4",
        "team": "System Enhancement",
        "pillar": "system",
        "photo": "",
    },
    {
        "id": "9",
        "name": "Pracilla Ullas",
        "role": "System Enhancement",
        "manager_id": "4",
        "team": "System Enhancement",
        "pillar": "system",
        "photo": "",
    },
    {
        "id": "10",
        "name": "Ayushi Choudhary",
        "role": "System Enhancement",
        "manager_id": "4",
        "team": "System Enhancement",
        "pillar": "system",
        "photo": "",
    },
    {
        "id": "11",
        "name": "Oktawia Marynowska",
        "role": "System Enhancement",
        "manager_id": "4",
        "team": "System Enhancement",
        "pillar": "system",
        "photo": "",
    },
    # Process Improvement team
    {
        "id": "12",
        "name": "Priyanka Kapoor",
        "role": "Process Improvement",
        "manager_id": "5",
        "team": "Process Improvement",
        "pillar": "process",
        "photo": "",
    },
    {
        "id": "13",
        "name": "Nehal Sharma",
        "role": "Process Improvement",
        "manager_id": "5",
        "team": "Process Improvement",
        "pillar": "process",
        "photo": "",
    },
    {
        "id": "14",
        "name": "Dinesh Sanmugam",
        "role": "Process Improvement",
        "manager_id": "5",
        "team": "Process Improvement",
        "pillar": "process",
        "photo": "",
    },
    # Automation team
    {
        "id": "15",
        "name": "Jasno John",
        "role": "Automation",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    {
        "id": "16",
        "name": "Meghana Gowda",
        "role": "Automation, BI & Reporting Jr.",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    {
        "id": "17",
        "name": "Deepak M.",
        "role": "Automation",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    {
        "id": "18",
        "name": "Shivashankar V.",
        "role": "Automation",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    {
        "id": "19",
        "name": "Anshul Chhabra",
        "role": "Automation",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    {
        "id": "20",
        "name": "Mehak Verma",
        "role": "Automation",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    {
        "id": "21",
        "name": "Dhanalakshmi A.",
        "role": "Automation",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
    {
        "id": "22",
        "name": "Ashutosh Uniyal",
        "role": "Automation",
        "manager_id": "6",
        "team": "Automation",
        "pillar": "automation",
        "photo": "",
    },
]

PILLARS = [
    {"id": "system", "label": "System Enhancement", "tone": "green"},
    {"id": "process", "label": "Process Improvement", "tone": "purple"},
    {"id": "automation", "label": "Automation", "tone": "orange"},
]

COLUMNS = ["id", "name", "role", "manager_id", "team", "pillar", "photo"]


def write_excel(people: list[dict]) -> Path:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PHOTOS_DIR.mkdir(parents=True, exist_ok=True)
    path = DATA_DIR / "org-chart.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Org Chart"

    header_fill = PatternFill("solid", fgColor="00008F")
    header_font = Font(color="FFFFFF", bold=True, name="Arial", size=11)
    note_font = Font(color="666666", italic=True, name="Arial", size=10)

    for col, key in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, col, key)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, person in enumerate(people, start=2):
        for col, key in enumerate(COLUMNS, start=1):
            ws.cell(row, col, person.get(key, ""))

    info = wb.create_sheet("How to edit", 1)
    lines = [
        "How to update the organisation chart",
        "",
        "1. Edit names, roles, teams, or pillars on the 'Org Chart' sheet.",
        "2. To add a person: add a new row with a unique id.",
        "3. Set manager_id to the id of the person they report to.",
        "4. Leave manager_id blank for the top of the chart (only one root).",
        "5. pillar must be one of: leadership | system | process | automation",
        "6. Photos:",
        "   a) Put image files in assets/photos/ (jpg/png/webp).",
        "   b) In the photo column, enter the filename only, e.g. rebecca.jpg",
        "      OR a full https:// image URL.",
        "   c) Leave photo blank to keep the square placeholder.",
        "7. Commit and push. GitHub Pages rebuilds the chart.",
        "",
        "Columns:",
        "id          Unique identifier (text or number).",
        "name        Display name on the chart card.",
        "role        Job title / role line under the name.",
        "manager_id  id of the manager; blank for the root.",
        "team        Team / department label.",
        "pillar      leadership | system | process | automation",
        "photo       Filename in assets/photos/ OR image URL.",
    ]
    for i, line in enumerate(lines, start=1):
        cell = info.cell(i, 1, line)
        if i == 1:
            cell.font = Font(bold=True, color="00008F", size=14, name="Arial")
        else:
            cell.font = note_font
    info.column_dimensions["A"].width = 92

    widths = {
        "id": 8,
        "name": 28,
        "role": 30,
        "manager_id": 12,
        "team": 22,
        "pillar": 12,
        "photo": 28,
    }
    for col, key in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = widths[key]

    wb.save(path)
    return path


def write_json(people: list[dict]) -> Path:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path = DATA_DIR / "org-chart.json"
    payload = {
        "title": "Continuous Improvement Initiatives",
        "subtitle": "ESS Claims Organization Chart",
        "footer": "Job Rotation — Meghana Gowda",
        "layout": "pillars",
        "pillars": PILLARS,
        "people": people,
    }
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    return path


def people_from_excel(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Org Chart"] if "Org Chart" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    people = []
    for row in rows[1:]:
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        person = {}
        for i, key in enumerate(headers):
            if not key:
                continue
            value = row[i] if i < len(row) else ""
            person[key] = "" if value is None else str(value).strip()
        if person.get("id") and person.get("name"):
            people.append(person)
    return people


def main() -> None:
    PHOTOS_DIR.mkdir(parents=True, exist_ok=True)
    excel_path = DATA_DIR / "org-chart.xlsx"
    # Prefer canonical roster when replacing placeholder demo data.
    if excel_path.exists():
        existing = people_from_excel(excel_path)
        names = {p.get("name", "") for p in existing}
        if "Paulina Strozyk-Kozlowska" in names:
            people = existing
        else:
            people = PEOPLE
    else:
        people = PEOPLE

    for person in people:
        for key in COLUMNS:
            person.setdefault(key, "")

    write_excel(people)
    write_json(people)
    print(f"Wrote {len(people)} people to data/org-chart.xlsx and data/org-chart.json")


if __name__ == "__main__":
    main()
